#!/usr/bin/env python3
# md_to_docx.py
# Markdown 파일을 Word 문서로 변환

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import re

def create_docx_from_md(md_file_path, docx_file_path):
    """Markdown 파일을 읽어서 Word 문서로 변환"""
    
    # Word 문서 생성
    doc = Document()
    
    # 스타일 설정
    styles = doc.styles
    
    # 제목 스타일 설정
    title_style = styles['Title']
    title_style.font.size = Pt(18)
    title_style.font.bold = True
    
    # 제목 1 스타일 설정
    heading1_style = styles['Heading 1']
    heading1_style.font.size = Pt(16)
    heading1_style.font.bold = True
    
    # 제목 2 스타일 설정
    heading2_style = styles['Heading 2']
    heading2_style.font.size = Pt(14)
    heading2_style.font.bold = True
    
    # 제목 3 스타일 설정
    heading3_style = styles['Heading 3']
    heading3_style.font.size = Pt(12)
    heading3_style.font.bold = True
    
    # 코드 스타일 설정
    code_style = styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = 'Courier New'
    code_style.font.size = Pt(10)
    
    # Markdown 파일 읽기
    with open(md_file_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # 줄 단위로 분리
    lines = content.split('\n')
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # 빈 줄 처리
        if not line:
            doc.add_paragraph()
            i += 1
            continue
        
        # 제목 처리
        if line.startswith('# '):
            # 메인 제목
            title = line[2:]
            doc.add_heading(title, level=0)
        elif line.startswith('## '):
            # 제목 1
            title = line[3:]
            doc.add_heading(title, level=1)
        elif line.startswith('### '):
            # 제목 2
            title = line[4:]
            doc.add_heading(title, level=2)
        elif line.startswith('#### '):
            # 제목 3
            title = line[5:]
            doc.add_heading(title, level=3)
        
        # 코드 블록 처리
        elif line.startswith('```'):
            # 코드 블록 시작
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            
            if code_lines:
                code_text = '\n'.join(code_lines)
                code_para = doc.add_paragraph(code_text, style='Code')
                code_para.paragraph_format.left_indent = Inches(0.5)
        
        # 인라인 코드 처리
        elif '`' in line:
            # 인라인 코드가 포함된 텍스트
            parts = line.split('`')
            para = doc.add_paragraph()
            for j, part in enumerate(parts):
                if j % 2 == 0:  # 일반 텍스트
                    if part:
                        para.add_run(part)
                else:  # 코드
                    if part:
                        code_run = para.add_run(part)
                        code_run.font.name = 'Courier New'
                        code_run.font.size = Pt(10)
        
        # 목록 처리
        elif line.startswith('- '):
            # 불릿 포인트
            item_text = line[2:]
            para = doc.add_paragraph(item_text, style='List Bullet')
        elif line.startswith('1. '):
            # 번호 목록
            item_text = line[3:]
            para = doc.add_paragraph(item_text, style='List Number')
        
        # 굵은 텍스트 처리
        elif '**' in line:
            # 굵은 텍스트가 포함된 텍스트
            parts = line.split('**')
            para = doc.add_paragraph()
            for j, part in enumerate(parts):
                if j % 2 == 0:  # 일반 텍스트
                    if part:
                        para.add_run(part)
                else:  # 굵은 텍스트
                    if part:
                        bold_run = para.add_run(part)
                        bold_run.bold = True
        
        # 일반 텍스트 처리
        else:
            doc.add_paragraph(line)
        
        i += 1
    
    # 문서 저장
    doc.save(docx_file_path)
    print(f"변환 완료: {docx_file_path}")

if __name__ == "__main__":
    # 파일 변환
    create_docx_from_md('프로젝트_현황_보고서.md', '프로젝트_현황_보고서.docx')
